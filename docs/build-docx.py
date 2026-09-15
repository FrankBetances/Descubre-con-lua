# -*- coding: utf-8 -*-
"""Genera docs/Descubre-con-Lua-Manual-Casos-de-Uso.docx A PARTIR del HTML.

    pip install python-docx lxml
    python3 docs/build-docx.py

Portado del constructor del proyecto anterior de la casa, incluida su lección:
allí el script llevaba el texto del manual DUPLICADO dentro, el HTML avanzó y el
Word se quedó describiendo una versión anterior sin que nada avisara.

Aquí hay una sola fuente, docs/manual-casos-de-uso.html, y de ella salen los tres
formatos: HTML, PDF (docs/build-pdf.js) y este Word. Lo que no se puede
representar en Word —los degradados, el redondeo de las tarjetas— se sustituye
por su equivalente sobrio; el texto es el mismo, siempre.
"""
import os
import re

import lxml.html
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCS = os.path.join(ROOT, 'docs')
SRC = os.path.join(DOCS, 'manual-casos-de-uso.html')
OUT = os.path.join(DOCS, 'Descubre-con-Lua-Manual-Casos-de-Uso.docx')

# Paleta de Lúa: el azul marítimo de Vigo que usa la propia app.
PRIMARY = RGBColor(0x1B, 0x49, 0x65)
SLATE = RGBColor(0x3D, 0x5A, 0x80)
INK = RGBColor(0x2B, 0x2D, 0x42)
INK2 = RGBColor(0x4A, 0x4E, 0x69)
MUTED = RGBColor(0x6B, 0x72, 0x80)
WARN_INK = RGBColor(0x8A, 0x6D, 0x3B)
OK_INK = RGBColor(0x2A, 0x9D, 0x8F)

FILL_HEAD = 'E8F4F8'
FILL_ZEBRA = 'F9FBFC'
CALLOUT_FILL = {'': 'EEF6F9', 'warn': 'FFFBEB', 'ok': 'EAFAF2', 'violet': 'EEF2F7'}
CALLOUT_INK = {'': PRIMARY, 'warn': WARN_INK, 'ok': OK_INK, 'violet': SLATE}


# --------------------------------------------------------------------------- util
def shade(cell_or_par, hexcolor):
    el = OxmlElement('w:shd')
    el.set(qn('w:val'), 'clear')
    el.set(qn('w:fill'), hexcolor)
    target = cell_or_par._tc.get_or_add_tcPr() if hasattr(cell_or_par, '_tc') \
        else cell_or_par._p.get_or_add_pPr()
    target.append(el)


def no_borders(table):
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + edge)
        e.set(qn('w:val'), 'none')
        borders.append(e)
    # El esquema de OOXML fija el ORDEN de los hijos de <w:tblPr>: con un
    # `append` el elemento cae detrás de w:tblLook y el documento deja de
    # validar contra el XSD (Word puede ofrecer «reparar» el fichero).
    tbl_pr = table._tbl.tblPr
    tbl_pr.insert_element_before(
        borders,
        'w:shd', 'w:tblLayout', 'w:tblCellMar', 'w:tblLook',
        'w:tblCaption', 'w:tblDescription', 'w:tblPrChange',
    )


def write_stamp(kind):
    """Deja constancia de QUÉ versión del HTML produjo este fichero.

    Es la salvaguarda contra un fallo ya visto: que la fuente avance y el
    documento generado se quede atrás sin que nada avise.
    tools/check_manual_build.py compara este sello con el HTML.
    """
    import hashlib
    import json
    digest = hashlib.sha256(open(SRC, 'rb').read()).hexdigest()
    path = os.path.join(DOCS, '.manual-build-stamp.json')
    data = {}
    if os.path.exists(path):
        data = json.loads(open(path, encoding='utf-8').read())
    data[kind] = digest
    with open(path, 'w', encoding='utf-8') as fh:
        json.dump(data, fh, indent=1, sort_keys=True)
        fh.write('\n')


def classes(el):
    return (el.get('class') or '').split()


def clean(text):
    """Colapsa el espacio en blanco del HTML como haría el navegador."""
    return re.sub(r'\s+', ' ', text or '')


class Builder:
    def __init__(self, doc):
        self.doc = doc

    # ---------------------------------------------------------------- runs
    def runs(self, par, el, bold=False, italic=False, color=None, size=None, mono=False):
        """Vuelca el contenido en línea de `el` (texto + hijos + colas) en `par`."""
        def emit(txt, **kw):
            if not txt:
                return
            r = par.add_run(txt)
            r.bold = kw.get('bold', bold)
            r.italic = kw.get('italic', italic)
            c = kw.get('color', color)
            if c is not None:
                r.font.color.rgb = c
            s = kw.get('size', size)
            if s is not None:
                r.font.size = Pt(s)
            if kw.get('mono', mono):
                r.font.name = 'Consolas'

        emit(clean(el.text))
        for child in el:
            if not isinstance(child.tag, str):
                continue
            tag = child.tag
            cls = classes(child)
            if tag == 'br':
                par.add_run().add_break()
            elif tag in ('strong', 'b'):
                self.runs(par, child, bold=True, italic=italic, color=color, size=size)
            elif tag in ('em', 'i'):
                self.runs(par, child, bold=bold, italic=True, color=color, size=size)
            elif tag == 'code':
                self.runs(par, child, bold=bold, italic=italic, color=color,
                          size=size, mono=True)
            elif tag == 'span' and ('small' in cls or 'muted' in cls):
                self.runs(par, child, bold=bold, italic=italic,
                          color=MUTED if 'muted' in cls else color, size=8.5)
            else:
                self.runs(par, child, bold=bold, italic=italic, color=color, size=size)
            emit(clean(child.tail))

    def para(self, el, style=None, space_after=6, **kw):
        p = self.doc.add_paragraph(style=style)
        p.paragraph_format.space_after = Pt(space_after)
        self.runs(p, el, **kw)
        return p

    def text_para(self, txt, bold=False, size=None, color=None, space_after=6,
                  italic=False):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        r = p.add_run(txt)
        r.bold, r.italic = bold, italic
        if size:
            r.font.size = Pt(size)
        if color is not None:
            r.font.color.rgb = color
        return p

    # ------------------------------------------------------------- bloques
    def table(self, el):
        rows = el.xpath('./tr | ./tbody/tr | ./thead/tr')
        if not rows:
            return
        ncols = max(len(r.xpath('./td | ./th')) for r in rows)
        t = self.doc.add_table(rows=0, cols=ncols)
        t.style = 'Table Grid'
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, tr in enumerate(rows):
            cells = tr.xpath('./td | ./th')
            row = t.add_row()
            for j in range(ncols):
                cell = row.cells[j]
                cell.paragraphs[0].paragraph_format.space_after = Pt(2)
                if j >= len(cells):
                    continue
                header = cells[j].tag == 'th'
                self.runs(cell.paragraphs[0], cells[j], bold=header,
                          color=PRIMARY if header else None, size=9.5)
                if header:
                    shade(cell, FILL_HEAD)
                elif i % 2 == 0:
                    shade(cell, FILL_ZEBRA)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def callout(self, el):
        variant = next((c for c in classes(el) if c in ('warn', 'ok', 'violet')), '')
        t = self.doc.add_table(rows=1, cols=1)
        cell = t.rows[0].cells[0]
        shade(cell, CALLOUT_FILL[variant])
        no_borders(t)
        first = True
        for child in el:
            if not isinstance(child.tag, str):
                continue
            p = cell.paragraphs[0] if first else cell.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            if 'lbl' in classes(child):
                self.runs(p, child, bold=True, color=CALLOUT_INK[variant], size=9)
            else:
                self.runs(p, child, size=10)
            first = False
        self.doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def uc(self, el):
        head = next((c for c in el if 'uc-head' in classes(c)), None)
        if head is not None:
            slate = 'violet' in classes(head)
            t = self.doc.add_table(rows=1, cols=1)
            cell = t.rows[0].cells[0]
            shade(cell, '3D5A80' if slate else '1B4965')
            no_borders(t)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            for div in head:
                if 'id' in classes(div):
                    self.runs(p, div, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), size=9)
                elif 'title' in classes(div):
                    p2 = cell.add_paragraph()
                    p2.paragraph_format.space_after = Pt(0)
                    self.runs(p2, div, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), size=13)
            self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        body = next((c for c in el if 'uc-body' in classes(c)), None)
        if body is not None:
            self.walk(body)

    def uc_grid(self, el):
        pairs = []
        pending = None
        for child in el:
            cls = classes(child)
            if 'k' in cls:
                pending = child
            elif 'v' in cls and pending is not None:
                pairs.append((pending, child))
                pending = None
        if not pairs:
            return
        t = self.doc.add_table(rows=0, cols=2)
        t.style = 'Table Grid'
        for k, v in pairs:
            row = t.add_row()
            for cell, el2, bold, color in ((row.cells[0], k, True, MUTED),
                                           (row.cells[1], v, False, None)):
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(1)
                self.runs(p, el2, bold=bold, color=color, size=9)
            shade(row.cells[0], FILL_ZEBRA)
        t.columns[0].width = Cm(4.2)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def flowmap(self, el):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        parts = [clean(n.text_content()).strip() for n in el if 'node' in classes(n)]
        r = p.add_run('  →  '.join(parts))
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = PRIMARY

    def toc(self, el):
        for li in el:
            if 'section-label' in classes(li):
                self.text_para(clean(li.text_content()).strip().upper(),
                               bold=True, size=9, color=PRIMARY, space_after=2)
                continue
            num = li.find('./span[@class="num"]')
            txt = li.find('./span[@class="txt"]')
            p = self.doc.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.left_indent = Cm(0.6)
            if num is not None:
                r = p.add_run(clean(num.text_content()).strip().ljust(9))
                r.bold = True
                r.font.color.rgb = PRIMARY
            if txt is not None:
                p.add_run(clean(txt.text_content()).strip())

    def cover(self, el):
        def first(cls):
            found = el.xpath('.//*[contains(concat(" ", @class, " "), " %s ")]' % cls)
            return found[0] if found else None

        for cls, size, color, bold, space in (('brand', 10, PRIMARY, True, 10),
                                              ('title', 26, PRIMARY, True, 6),
                                              ('subtitle', 13, INK2, False, 22)):
            node = first(cls)
            if node is None:
                continue
            p = self.doc.add_paragraph()
            p.paragraph_format.space_after = Pt(space)
            self.runs(p, node, bold=bold, color=color, size=size)

        for node in el.xpath('.//p[contains(@class, "meta")]'):
            self.para(node, size=9.5, color=MUTED, space_after=8)

        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # -------------------------------------------------------------- recorrido
    def shots(self, el):
        """Las capturas, en la misma pareja gallego/castellano que el HTML.

        Word no tiene flex: cada par va en una tabla de dos columnas sin
        bordes, que es lo que más se le parece. Sin esto las imágenes se
        perdían EN SILENCIO —el constructor no miraba los `img`— y el Word
        habría pasado el gate del manual describiendo un documento con
        capturas sin llevar ninguna.
        """
        for pair in el.xpath('.//div[contains(@class, "pair")]'):
            figuras = pair.xpath('./figure')
            if not figuras:
                continue
            table = self.doc.add_table(rows=1, cols=len(figuras))
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            no_borders(table)
            for i, fig in enumerate(figuras):
                cell = table.cell(0, i)
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                src = (fig.xpath('./img/@src') or [None])[0]
                ruta = os.path.join(DOCS, src) if src else None
                if ruta and os.path.exists(ruta):
                    cell.paragraphs[0].add_run().add_picture(ruta, width=Cm(6.4))
                else:
                    # Nunca en silencio: si falta el PNG hay que verlo.
                    raise SystemExit(
                        'build-docx: falta la captura %s. El Word no puede '
                        'salir sin ella.' % (src or '(sin src)'))
                pie = fig.xpath('./figcaption')
                if pie:
                    par = cell.add_paragraph()
                    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    par.paragraph_format.space_after = Pt(8)
                    run = par.add_run(clean(pie[0].text_content()).strip())
                    run.font.size = Pt(8.5)
                    run.font.color.rgb = MUTED
            self.doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def walk(self, root):
        for el in root:
            if not isinstance(el.tag, str):
                continue  # comentarios
            cls = classes(el)
            tag = el.tag

            if 'cover' in cls:
                self.cover(el)
            elif 'chapter-kicker' in cls:
                self.text_para(clean(el.text_content()).strip().upper(),
                               bold=True, size=8.5, color=PRIMARY, space_after=1)
            elif tag == 'h2':
                self.doc.add_heading(clean(el.text_content()).strip(), level=1)
            elif tag == 'h3':
                self.doc.add_heading(clean(el.text_content()).strip(), level=2)
            elif tag == 'h4':
                self.doc.add_heading(clean(el.text_content()).strip(), level=3)
            elif 'toc' in cls:
                self.toc(el)
            elif 'shots' in cls:
                self.shots(el)
            elif 'flowmap' in cls:
                self.flowmap(el)
            elif 'callout' in cls:
                self.callout(el)
            elif 'uc-grid' in cls:
                self.uc_grid(el)
            elif 'uc' in cls:
                self.uc(el)
            elif 'footer-note' in cls:
                self.para(el, size=8.5, color=MUTED, italic=True, space_after=2)
            elif tag == 'table':
                self.table(el)
            elif tag == 'p':
                small = 'small' in cls
                self.para(el, size=9 if small else None,
                          color=MUTED if 'muted' in cls else None)
            elif tag in ('ul', 'ol'):
                style = 'List Number' if tag == 'ol' else 'List Bullet'
                for li in el:
                    if not isinstance(li.tag, str):
                        continue
                    nested = [c for c in li if isinstance(c.tag, str)
                              and c.tag in ('ul', 'ol')]
                    p = self.doc.add_paragraph(style=style)
                    p.paragraph_format.space_after = Pt(2)
                    self.runs(p, li)
                    for sub in nested:
                        for subli in sub:
                            sp = self.doc.add_paragraph(style=style)
                            sp.paragraph_format.left_indent = Cm(1.6)
                            sp.paragraph_format.space_after = Pt(1)
                            self.runs(sp, subli)
            elif tag == 'section':
                if 'page-break' in cls:
                    self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                self.walk(el)
            elif tag in ('div', 'figure'):
                self.walk(el)


def build():
    doc = Document()
    doc.core_properties.title = 'Descubre con Lúa · Manual de Casos de Uso'
    doc.core_properties.author = 'Descubre con Lúa · Edición Vigo'
    doc.core_properties.language = 'es-ES'

    # La plantilla de python-docx trae <w:zoom> sin el atributo w:percent, que
    # el esquema exige. Sale en la validación XSD, así que se completa aquí.
    zoom = doc.settings.element.find(qn('w:zoom'))
    if zoom is not None and zoom.get(qn('w:percent')) is None:
        zoom.set(qn('w:percent'), '100')

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.2)

    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')

    for name, size, color, before in (('Heading 1', 17, PRIMARY, 16),
                                      ('Heading 2', 13, INK, 12),
                                      ('Heading 3', 11.5, INK2, 10)):
        st = doc.styles[name]
        st.font.name = 'Calibri'
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(5)

    html = lxml.html.parse(SRC).getroot()
    b = Builder(doc)
    b.walk(html.find('body'))
    doc.save(OUT)
    write_stamp('docx')
    print('escrito %s' % os.path.relpath(OUT, ROOT))


if __name__ == '__main__':
    build()
